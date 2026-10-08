"""투자 판단 시트(.xlsx)를 만든다. 구글 시트로 가져가도, 엑셀로 열어도 같은 수식으로 계산된다.

모든 판단은 수식으로 계산한다. 파이썬은 시트를 만들기만 하고 결과를 써 넣지 않는다.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

from judge.common import BANDS, IMPL, SPEC
from judge.macro import SERIES as MACRO_SERIES

FONT = "Arial"
F_INPUT = PatternFill("solid", fgColor="FFF2CC")
F_HEAD = PatternFill("solid", fgColor="D9D9D9")
F_GROUP = PatternFill("solid", fgColor="EDEDED")
F_KEY = PatternFill("solid", fgColor="DDEBF7")
F_RED = PatternFill("solid", fgColor="F8CBAD")
F_ORANGE = PatternFill("solid", fgColor="FCE4D6")
F_YELLOW = PatternFill("solid", fgColor="FFF2CC")
F_BLUE = PatternFill("solid", fgColor="DDEBF7")
F_GREEN = PatternFill("solid", fgColor="E2EFDA")

FMT_DATE = "yyyy-mm-dd"
FMT_INT = "#,##0"
FMT_DEC = "#,##0.00"
FMT_PCT = "0.0%"
FMT_SCORE = "0.0"

# 표 범위
STOCK_FIRST, STOCK_LAST = 6, 105
HOLD_FIRST, HOLD_LAST = 11, 20  # 명세: 개별 종목 3~5개
HOLD_ACC, HOLD_CASH, HOLD_FX, HOLD_N, HOLD_SUM = "$D$3", "$D$4", "$D$5", "$D$6", "$D$7"
REASON_FIRST, REASON_LAST = 6, 205
PLAN_FIRST, PLAN_LAST = 6, 25
RAW_FIRST = 3  # 원자료 첫 줄(2행은 머리글)
RAW_DAILY, RAW_MONTHLY = 4000, 600  # 원자료 칸 수
SENT_HDR, SENT_FIRST = 18, 19
CAL_ROWS = IMPL["calendar_rows"]
SENT_LAST = SENT_FIRST + CAL_ROWS - 1
SENT_HELP_OFFSET = SENT_FIRST - RAW_FIRST  # 원자료 i행 ↔ 보조열 i+16행
MACRO_HELP_FIRST = 21
MACRO_HELP_OFFSET = MACRO_HELP_FIRST - RAW_FIRST

SH_GUIDE, SH_DASH, SH_SET = "안내", "대시보드", "설정"
SH_STOCK = {"한국": "종목_한국", "미국": "종목_미국"}
SH_HOLD, SH_REASON, SH_PLAN, SH_LOG = "보유", "매수이유", "매수계획", "판단기록"
SH_SENT = {"미국": "심리_미국", "한국": "심리_한국"}
SH_RAW = {"미국": "자료_미국심리", "한국": "자료_한국심리"}
SH_MACRO, SH_MACRO_RAW = "거시", "자료_거시"


def ref(sheet: str, cell: str) -> str:
    return f"'{sheet}'!{cell}"


def rng(sheet: str, c: str, r1: int, r2: int) -> str:
    return f"'{sheet}'!${c}${r1}:${c}${r2}"


def pct_expr(x: str, rng_: str, n: str) -> str:
    """풀 안 백분위. 평균 순위, 하나뿐이면 50."""
    return f"IF({n}=1,50,(_xlfn.RANK.AVG({x},{rng_},1)-1)/({n}-1)*100)"


# 설정 탭 칸
def _s(cell):
    c, r = cell[0], cell[1:]
    return f"'{SH_SET}'!${c}${r}"


CFG = {
    "cand": _s("B5"),
    "repl_below": _s("B6"),
    "repl_gap": _s("B7"),
    "debt_max": _s("B8"),
    "mcap": {"한국": _s("B9"), "미국": _s("B10")},
    "loss_acc": _s("B11"),
    "loss_price": _s("B12"),
    "conc": _s("B13"),
    "conc_n": _s("B14"),
    "sent_m": _s("B15"),
    "macro_m": _s("B16"),
    "chg_m": _s("B17"),
    "strong": _s("B20"),
    "weak": _s("B21"),
    "stale": _s("B22"),
    "slack": _s("B23"),
}
BAND_FIRST = 27
BAND_LAST = BAND_FIRST + len(BANDS) - 1
BAND = {k: f"'{SH_SET}'!${c}${BAND_FIRST}:${c}${BAND_LAST}" for k, c in zip(("lo", "name", "n", "w"), "ABCD")}


def band_lookup(score: str, what: str) -> str:
    return f"INDEX({BAND[what]},MATCH({score},{BAND['lo']},1))"


# ---------------------------------------------------------------- 서식 도구


class W:
    """시트 쓰기 도우미."""

    def __init__(self, ws):
        self.ws = ws

    def put(self, cell, value, *, fmt=None, bold=False, fill=None, color=None, size=None, wrap=False, italic=False):
        c = self.ws[cell]
        c.value = value
        if fmt:
            c.number_format = fmt
        if bold or color or size or italic:
            c.font = Font(name=FONT, bold=bold, color=color, size=size or 10, italic=italic)
        if fill:
            c.fill = fill
        if wrap:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        return c

    def title(self, text, sub=None):
        self.put("A1", text, bold=True, size=14)
        if sub:
            self.put("A2", sub, italic=True, color="595959")

    def header(self, row, labels, start=1, fill=F_HEAD):
        for i, label in enumerate(labels):
            c = self.ws.cell(row, start + i, label)
            c.font = Font(name=FONT, bold=True, size=10)
            c.fill = fill
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    def input_cells(self, cols, r1, r2, fmt=None):
        for c in cols:
            for r in range(r1, r2 + 1):
                cell = self.ws[f"{c}{r}"]
                cell.fill = F_INPUT
                cell.font = Font(name=FONT, color="0000FF", size=10)
                if fmt:
                    cell.number_format = fmt

    def fmt(self, cols, r1, r2, fmt):
        for c in cols:
            for r in range(r1, r2 + 1):
                self.ws[f"{c}{r}"].number_format = fmt

    def widths(self, mapping):
        for c, w in mapping.items():
            self.ws.column_dimensions[c].width = w

    def hide(self, first, last):
        for i in range(first, last + 1):
            self.ws.column_dimensions[col(i)].hidden = True

    def validate(self, rng_, options, prompt=None):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        if prompt:
            dv.prompt, dv.showInputMessage = prompt, True
        self.ws.add_data_validation(dv)
        dv.add(rng_)

    def note(self, cell, text):
        self.ws[cell].comment = Comment(text, "투자판단")


def ci(letter: str) -> int:
    from openpyxl.utils import column_index_from_string

    return column_index_from_string(letter)


# ---------------------------------------------------------------- 설정


def build_settings(ws):
    w = W(ws)
    w.title("설정", "판단 기준 값. 명세 확정값은 바꾸지 않는다. 수식이 모두 이 칸을 읽는다.")
    w.header(4, ["항목", "값", "단위", "근거"])
    rows = [
        ("편입 후보: 점수 이상", SPEC["candidate_score"], "점", "명세 1", None),
        ("교체: 보유 종목 점수 미만", SPEC["replace_below"], "점", "명세 1·6", None),
        ("교체: 후보와 점수 차 이상", SPEC["replace_gap"], "점", "명세 1·6", None),
        ("부채비율(총부채 ÷ 총자본) 상한", SPEC["debt_ratio_max"], "", "명세 2", "0%"),
        ("시가총액 하한 — 한국", SPEC["mcap_min"]["한국"], "억 원", "명세 2", FMT_INT),
        ("시가총액 하한 — 미국", SPEC["mcap_min"]["미국"], "백만 달러", "명세 2", FMT_INT),
        ("손실 한도: 계좌 평가액 대비 손실액", SPEC["loss_account"], "", "명세 5", FMT_PCT),
        ("손실 한도: 매수 평균가 대비 가격", SPEC["loss_price"], "", "명세 5", FMT_PCT),
        ("집중 한도: 개별 몫(주식+대기 현금) 대비", SPEC["conc_max"], "", "명세 6", FMT_PCT),
        ("집중 규칙 적용: 보유 종목 수 이상", SPEC["conc_min_holdings"], "종목", "명세 6", None),
        ("심리 지수 기간", SPEC["sentiment_months"], "개월", "명세 3 (5년)", None),
        ("거시 패널 기간", SPEC["macro_months"], "개월", "명세 4 (10년)", None),
        ("거시 변화 기간", SPEC["change_months"], "개월", "명세 4", None),
    ]
    for i, (label, value, unit, why, fmt) in enumerate(rows, start=5):
        w.put(f"A{i}", label)
        w.put(f"B{i}", value, fmt=fmt, bold=True)
        w.put(f"C{i}", unit)
        w.put(f"D{i}", why)
    w.put("A19", "구현하는 쪽이 정한 값", bold=True)
    impl = [
        ("한 줄 요약: 강점 백분위 이상", IMPL["strong_pct"], "백분위", "편입 기준(상위 약 30%)과 맞춤"),
        ("한 줄 요약: 약점 백분위 미만", IMPL["weak_pct"], "백분위", "하위 약 30%"),
        ("심리 지표 자료 지연 경고", IMPL["stale_days"], "일", "기준일보다 이만큼 넘게 오래되면 경고"),
        ("거시 10년 자료 부족 판정 여유", IMPL["coverage_slack_days"], "일", "창 첫 자료가 이만큼 늦으면 경고"),
    ]
    for i, (label, value, unit, why) in enumerate(impl, start=20):
        w.put(f"A{i}", label)
        w.put(f"B{i}", value, bold=True)
        w.put(f"C{i}", unit)
        w.put(f"D{i}", why)
    w.put(f"A{BAND_FIRST - 2}", "심리 구간 (하한 이상 ~ 다음 하한 미만) — 명세 3", bold=True)
    w.header(BAND_FIRST - 1, ["점수 하한", "구간", "신규 매수 횟수", "간격(주)"])
    for i, (lo, name, n, wk) in enumerate(BANDS, start=BAND_FIRST):
        w.put(f"A{i}", lo, bold=True)
        w.put(f"B{i}", name)
        w.put(f"C{i}", n)
        w.put(f"D{i}", wk)
    w.put("A33", "수식에 고정된 값 (바꾸려면 시트를 다시 만들어야 함)", bold=True)
    fixed = [
        f"이동평균 {SPEC['ma_days']}거래일, 신용융자 증감 {SPEC['credit_lag']}거래일 (명세 3)",
        f"심리 계산은 지수의 최근 {CAL_ROWS:,}거래일까지 사용 (5년 + 125일을 덮음)",
        f"원자료 칸: 일별 {RAW_DAILY:,}행, 월별 {RAW_MONTHLY:,}행",
        "평균값(종목 지표 평균, 심리 지표 평균)은 소수 6자리 반올림 뒤 순위. 동점은 평균 순위",
    ]
    for i, t in enumerate(fixed, start=34):
        w.put(f"A{i}", "· " + t)
    w.widths({"A": 42, "B": 12, "C": 14, "D": 36})


# ---------------------------------------------------------------- 종목


# 종목 시트 열
SC = {
    "code": "A", "name": "B", "held": "C", "fin": "D", "basis": "E",
    "ni": "F", "eq": "G", "liab": "H", "op": "I", "rev": "J", "rev3": "K",
    "mcap": "L", "debt": "M", "cash": "N", "p12": "O", "p1": "P",
    "decision": "Q", "score": "R", "summary": "S",
    "c_roe": "T", "c_growth": "U", "c_ey": "V", "c_mom": "W",
    "filter": "X", "debt_ratio": "Y",
    "roe": "Z", "growth": "AA", "ey": "AB", "mom": "AC",
    "missing": "AD", "status": "AE", "scored": "AF",
    "w_roe": "AG", "w_growth": "AH", "w_ey": "AI", "w_mom": "AJ",
    "p_roe": "AK", "p_growth": "AL", "p_ey": "AM", "p_mom": "AN",
    "avg": "AO", "cand": "AP", "rankkey": "AQ", "reasons": "AR",
}
STOCK_POOL_N = "$B$3"
STOCK_BEST_NAME = "$E$3"
STOCK_BEST_SCORE = "$H$3"
STOCK_PRICE_DATE = "$K$3"
METRIC_KEYS = ("roe", "growth", "ey", "mom")
PHRASES = {
    "roe": ("수익성 높음", "수익성 낮음", "수익성"),
    "growth": ("성장 빠름", "성장 느림", "성장"),
    "ey": ("가격 쌈", "가격 비쌈", "가격"),
    "mom": ("추세 좋음", "추세 약함", "추세"),
}


def build_stock(ws, country: str):
    w = W(ws)
    unit = "억 원" if country == "한국" else "백만 달러"
    cur = "원" if country == "한국" else "달러"
    w.title(
        f"종목 점수 — {country}",
        f"금액 단위: {unit} · 주가: {cur}(수정주가) · 보유 종목과 후보를 모두 적는다. 순위는 이 시트(같은 국가) 안에서만 매긴다.",
    )
    w.put("A3", "채점 종목 수", bold=True)
    w.put("D3", "최고 미보유 후보", bold=True)
    w.put("G3", "점수", bold=True)
    w.put("J3", "주가 기준일", bold=True)
    w.input_cells(["K"], 3, 3, FMT_DATE)
    w.note("K3", "12개월 전·1개월 전 주가를 잡은 기준일(메모). 모든 종목이 같은 기준일이어야 한다.")
    first, last = STOCK_FIRST, STOCK_LAST

    def R(c):
        return f"${c}${first}:${c}${last}"

    w.put("B3", f"=COUNT({R(SC['scored'])})")
    w.put("H3", f'=IF(COUNT({R(SC["cand"])})=0,"",MAX({R(SC["cand"])}))', fmt=FMT_SCORE, bold=True)
    w.put("E3", f'=IF(H3="","",INDEX({R(SC["name"])},MATCH(H3,{R(SC["cand"])},0))&"")', bold=True)

    ws.merge_cells("A4:P4")
    ws.merge_cells("Q4:Y4")
    w.put("A4", "입력 (노란 칸)", bold=True, fill=F_GROUP)
    w.put("Q4", "결과 — 점수 설명 카드", bold=True, fill=F_KEY)
    labels = [
        "종목코드", "종목명", "보유\n(Y)", "금융·리츠\n(Y)", "재무 기준\n(메모)",
        "순이익\n최근 1년", "자본총계", "부채총계", "영업이익\n최근 1년", "매출\n최근 1년", "매출\n3년 전 1년",
        "시가총액", "이자부부채", "현금성자산", "주가\n12개월 전", "주가\n1개월 전",
        "판정", "점수\n(0~100)", "한 줄 요약",
        "ROE\n값 · 풀 안 위치", "매출성장률(3년 연평균)\n값 · 풀 안 위치", "이익수익률\n값 · 풀 안 위치", "12-1개월 수익률\n값 · 풀 안 위치",
        "거름망", "부채비율",
        "ROE", "매출성장률", "이익수익률", "12-1 수익률", "결측", "상태", "채점(1)",
        "ROE(채점)", "성장(채점)", "이익수익률(채점)", "12-1(채점)",
        "ROE 백분위", "성장 백분위", "이익수익률 백분위", "12-1 백분위",
        "평균 백분위", "미보유 후보 점수", "대시보드 순위값", "거름망 사유",
    ]
    w.header(5, labels)
    ws.row_dimensions[5].height = 42
    st, wk = CFG["strong"], CFG["weak"]
    mcap_min = CFG["mcap"][country]

    for r in range(first, last + 1):
        c = {k: f"{v}{r}" for k, v in SC.items()}
        f = {}
        f["debt_ratio"] = f'=IF(AND(ISNUMBER({c["eq"]}),ISNUMBER({c["liab"]})),IF({c["eq"]}>0,{c["liab"]}/{c["eq"]},""),"")'
        f["reasons"] = (
            f'=IF({c["code"]}="","",_xlfn.TEXTJOIN(", ",TRUE,'
            f'IF({c["fin"]}="Y","금융·리츠",""),'
            f'IF(NOT(ISNUMBER({c["op"]})),"영업이익 없음",IF({c["op"]}<=0,"영업적자","")),'
            f'IF(OR(NOT(ISNUMBER({c["eq"]})),NOT(ISNUMBER({c["liab"]}))),"부채비율 자료 없음",'
            f'IF({c["eq"]}<=0,"자본잠식",IF({c["liab"]}/{c["eq"]}>{CFG["debt_max"]},"부채비율 초과",""))),'
            f'IF(NOT(ISNUMBER({c["mcap"]})),"시총 없음",IF({c["mcap"]}<{mcap_min},"시총 미달",""))))'
        )
        f["filter"] = f'=IF({c["code"]}="","",IF({c["reasons"]}="","통과","제외: "&{c["reasons"]}))'
        f["roe"] = f'=IF(AND(ISNUMBER({c["ni"]}),ISNUMBER({c["eq"]})),IF({c["eq"]}>0,{c["ni"]}/{c["eq"]},""),"")'
        f["growth"] = (
            f'=IF(AND(ISNUMBER({c["rev"]}),ISNUMBER({c["rev3"]})),'
            f'IF(AND({c["rev"]}>0,{c["rev3"]}>0),({c["rev"]}/{c["rev3"]})^(1/3)-1,""),"")'
        )
        ev = f'({c["mcap"]}+{c["debt"]}-{c["cash"]})'
        f["ey"] = (
            f'=IF(AND(ISNUMBER({c["op"]}),ISNUMBER({c["mcap"]}),ISNUMBER({c["debt"]}),ISNUMBER({c["cash"]})),'
            f'IF({ev}>0,{c["op"]}/{ev},""),"")'
        )
        f["mom"] = (
            f'=IF(AND(ISNUMBER({c["p12"]}),ISNUMBER({c["p1"]})),'
            f'IF(AND({c["p12"]}>0,{c["p1"]}>0),{c["p1"]}/{c["p12"]}-1,""),"")'
        )
        f["missing"] = (
            f'=_xlfn.TEXTJOIN(", ",TRUE,IF({c["roe"]}="","ROE",""),IF({c["growth"]}="","매출성장률",""),'
            f'IF({c["ey"]}="","이익수익률",""),IF({c["mom"]}="","12-1 수익률",""))'
        )
        f["status"] = (
            f'=IF({c["code"]}="","",IF({c["reasons"]}<>"","점수 없음(거름망)",'
            f'IF({c["missing"]}="","채점","점수 없음(결측: "&{c["missing"]}&")")))'
        )
        f["scored"] = f'=IF({c["status"]}="채점",1,"")'
        for k in METRIC_KEYS:
            f[f"w_{k}"] = f'=IF({c["scored"]}=1,{c[k]},"")'
            f[f"p_{k}"] = f'=IF({c["scored"]}=1,{pct_expr(c["w_" + k], R(SC["w_" + k]), STOCK_POOL_N)},"")'
        f["avg"] = f'=IF({c["scored"]}=1,ROUND(AVERAGE({c["p_roe"]}:{c["p_mom"]}),{IMPL["tie_digits"]}),"")'
        f["score"] = f'=IF({c["scored"]}=1,ROUND({pct_expr(c["avg"], R(SC["avg"]), STOCK_POOL_N)},1),"")'
        f["cand"] = f'=IF(AND({c["scored"]}=1,{c["held"]}<>"Y"),{c["score"]},"")'

        def tier(kind):
            parts = []
            for k in METRIC_KEYS:
                p = c["p_" + k]
                strong, weak, mid = PHRASES[k]
                if kind == "strong":
                    parts.append(f'IF({p}>={st},"{strong}","")')
                elif kind == "weak":
                    parts.append(f'IF({p}<{wk},"{weak}","")')
                else:
                    parts.append(f'IF(AND({p}>={wk},{p}<{st}),"{mid}","")')
            return ",".join(parts)

        mid_join = f'_xlfn.TEXTJOIN("·",TRUE,{tier("mid")})'
        summary = (
            f'_xlfn.TEXTJOIN(". ",TRUE,_xlfn.TEXTJOIN(", ",TRUE,{tier("strong")}),'
            f'_xlfn.TEXTJOIN(", ",TRUE,{tier("weak")}),IF({mid_join}="","","보통: "&{mid_join}))'
        )
        f["summary"] = (
            f'=IF({c["code"]}="","",IF({c["scored"]}=1,{summary},'
            f'IF({c["filter"]}="통과",{c["status"]},{c["filter"]})))'
        )
        for k in METRIC_KEYS:
            top = f'({STOCK_POOL_N}-_xlfn.RANK.AVG({c["w_" + k]},{R(SC["w_" + k])},1)+1)/{STOCK_POOL_N}*100'
            f["c_" + k] = (
                f'=IF({c["scored"]}=1,TEXT({c[k]},"0.0%")&" · 상위 "&TEXT({top},"0")&"%",'
                f'IF({c["code"]}="","",IF({c[k]}="","결측",TEXT({c[k]},"0.0%"))))'
            )
        f["decision"] = (
            f'=IF({c["code"]}="","",IF({c["held"]}="Y",IF({c["scored"]}=1,'
            f'IF(AND({c["score"]}<{CFG["repl_below"]},ISNUMBER({STOCK_BEST_SCORE})),'
            f'IF(ROUND({STOCK_BEST_SCORE}-{c["score"]},1)>={CFG["repl_gap"]},'
            f'"교체 제안 → "&{STOCK_BEST_NAME}&" ("&TEXT({STOCK_BEST_SCORE},"0.0")&")","보유 유지"),"보유 유지"),'
            f'"보유 · 점수 없음"),IF({c["scored"]}=1,IF({c["score"]}>={CFG["cand"]},"편입 후보",""),"")))'
        )
        f["rankkey"] = f'=IF(AND(ISNUMBER({c["cand"]}),{c["decision"]}="편입 후보"),{c["cand"]}+(1000-ROW())/1000000,"")'
        for k, formula in f.items():
            ws[c[k]] = formula

    w.input_cells([SC[k] for k in ("code", "name", "held", "fin", "basis")], first, last)
    w.fmt([SC["code"]], first, last, "@")
    money = [SC[k] for k in ("ni", "eq", "liab", "op", "rev", "rev3", "mcap", "debt", "cash")]
    w.input_cells(money, first, last, FMT_INT)
    w.input_cells([SC["p12"], SC["p1"]], first, last, FMT_INT if country == "한국" else FMT_DEC)
    w.fmt([SC["score"]], first, last, FMT_SCORE)
    w.fmt([SC["debt_ratio"]], first, last, "0%")
    w.fmt([SC[k] for k in METRIC_KEYS], first, last, FMT_PCT)
    for r in range(first, last + 1):
        ws[f"{SC['decision']}{r}"].font = Font(name=FONT, bold=True, size=10)
        ws[f"{SC['score']}{r}"].font = Font(name=FONT, bold=True, size=10)
    w.validate(f"{SC['held']}{first}:{SC['fin']}{last}", ["Y"])
    dec = f"{SC['decision']}{first}:{SC['decision']}{last}"
    ws.conditional_formatting.add(dec, FormulaRule(formula=[f'{SC["decision"]}{first}="편입 후보"'], fill=F_GREEN))
    ws.conditional_formatting.add(dec, FormulaRule(formula=[f'LEFT({SC["decision"]}{first},2)="교체"'], fill=F_ORANGE))
    w.widths(
        {
            "A": 10, "B": 16, "C": 6, "D": 8, "E": 9, "F": 10, "G": 10, "H": 10, "I": 10, "J": 10, "K": 10,
            "L": 11, "M": 10, "N": 10, "O": 10, "P": 10, "Q": 30, "R": 8, "S": 40,
            "T": 18, "U": 18, "V": 18, "W": 18, "X": 26, "Y": 9,
        }
    )
    w.hide(ci("Z"), ci("AR"))
    ws.freeze_panes = "C6"
    notes = {
        "F5": "연결 기준 당기순이익, 최근 4개 분기 합(TTM). 최근 보고서가 사업보고서면 그 연간 값.",
        "G5": "연결 자본총계(최근 분기말). ROE의 분모이자 부채비율의 분모.",
        "H5": "연결 부채총계(최근 분기말).",
        "I5": "연결 영업이익, 최근 4개 분기 합.",
        "J5": "연결 매출, 최근 4개 분기 합.",
        "K5": "J와 같은 기간을 3년 앞당긴 1년 매출. 3년 연평균 성장률 = (J/K)^(1/3) - 1.",
        "L5": "현재 시가총액(보통주).",
        "M5": "이자를 내는 부채: 단기차입금 + 유동성장기부채 + 사채 + 장기차입금. 리스부채 제외. 없으면 0.",
        "N5": "현금및현금성자산 + 단기금융상품(단기투자자산). 없으면 0.",
        "O5": "주가 기준일로부터 12개월 전 종가(수정주가).",
        "P5": "주가 기준일로부터 1개월 전 종가(수정주가). 12-1 수익률 = P/O - 1.",
        "D5": "은행·보험·증권 등 금융업과 리츠면 Y. 후보에서 제외된다.",
        "C5": "보유 중이면 Y. 보유 시트에도 같은 종목코드로 적는다.",
    }
    for k, v in notes.items():
        w.note(k, v)


# ---------------------------------------------------------------- 매수 이유


def build_reasons(ws):
    w = W(ws)
    w.title(
        "매수 이유 기록 (명세 7)",
        "매수 전에 이유 1~3개, 그중 핵심 1개(Y). 이유마다 확인할 근거와 무너짐 기준을 함께 쓴다. 근거를 하나도 쓸 수 없으면 사지 않는다.",
    )
    w.put("A3", "예) 데이터센터 매출이 빠르게 는다 / 부문 매출 전년비 / 2회 연속 10% 미만 · 핵심 고객과 장기 계약 / 계약 유지 공시 / 해지·미갱신", italic=True, color="595959")
    w.header(5, ["종목코드", "종목명", "이유\n번호", "매수 이유", "핵심\n(Y)", "확인할 근거", "무너짐 기준", "최근\n점검일", "점검 결과", "메모", "안내", "미기재", "핵심 붕괴", "실적 발표일", "점검 필요"])
    ws.row_dimensions[5].height = 32
    hold_code = rng(SH_HOLD, "B", HOLD_FIRST, HOLD_LAST)
    hold_earn = rng(SH_HOLD, "H", HOLD_FIRST, HOLD_LAST)
    for r in range(REASON_FIRST, REASON_LAST + 1):
        ws[f"L{r}"] = f'=IF(AND(A{r}<>"",OR(F{r}="",G{r}="")),1,0)'
        ws[f"M{r}"] = f'=IF(AND(A{r}<>"",E{r}="Y",I{r}="무너짐"),1,0)'
        earn = f"INDEX({hold_earn},MATCH(A{r},{hold_code},0))"
        ws[f"N{r}"] = f'=IF(A{r}="","",IFERROR(IF({earn}="","",{earn}),""))'
        ws[f"O{r}"] = f'=IF(ISNUMBER(N{r}),IF(OR(NOT(ISNUMBER(H{r})),H{r}<N{r}),1,0),0)'
        ws[f"K{r}"] = (
            f'=IF(A{r}="","",IF(L{r}=1,"근거·무너짐 기준을 채우세요",IF(M{r}=1,"핵심 이유 무너짐 → 매도 검토",'
            f'IF(O{r}=1,"실적 발표 뒤 점검 필요",""))))'
        )
    w.input_cells(list("ABCDEFGHIJ"), REASON_FIRST, REASON_LAST)
    w.fmt(["A"], REASON_FIRST, REASON_LAST, "@")
    w.fmt(["H"], REASON_FIRST, REASON_LAST, FMT_DATE)
    w.validate(f"E{REASON_FIRST}:E{REASON_LAST}", ["Y"])
    w.validate(f"I{REASON_FIRST}:I{REASON_LAST}", ["유지", "무너짐"])
    w.widths({"A": 10, "B": 14, "C": 6, "D": 34, "E": 6, "F": 28, "G": 24, "H": 11, "I": 9, "J": 24, "K": 28})
    w.hide(ci("L"), ci("O"))
    ws.freeze_panes = "C6"


def record_status_expr(code: str) -> str:
    a = rng(SH_REASON, "A", REASON_FIRST, REASON_LAST)
    e = rng(SH_REASON, "E", REASON_FIRST, REASON_LAST)
    lcol = rng(SH_REASON, "L", REASON_FIRST, REASON_LAST)
    n = f"COUNTIF({a},{code})"
    return (
        f'IF({n}=0,"기록 없음",IF({n}>3,"이유 3개 초과",IF(COUNTIFS({a},{code},{e},"Y")<>1,"핵심 1개 표시 필요",'
        f'IF(SUMIFS({lcol},{a},{code})>0,"근거·기준 미기재","기록 완료"))))'
    )


def score_lookup(country_cell: str, code: str) -> str:
    def one(country):
        sh = SH_STOCK[country]
        return f'IFERROR(INDEX({rng(sh, SC["score"], STOCK_FIRST, STOCK_LAST)},MATCH({code},{rng(sh, "A", STOCK_FIRST, STOCK_LAST)},0)),"")'

    return f'IF({country_cell}="한국",{one("한국")},IF({country_cell}="미국",{one("미국")},""))'


# ---------------------------------------------------------------- 보유


def build_holdings(ws):
    w = W(ws)
    w.title("보유 종목 · 매도 판단 (명세 5·6)", "프로그램은 제안만 한다. '많이 올랐다'는 매도 신호가 아니다.")
    labels = [
        ("A3", "전체 투자 계좌 평가액(원)"),
        ("A4", "개별 몫 대기 현금(원)"),
        ("A5", "현재 원/달러 환율"),
        ("A6", "보유 개별 종목 수"),
        ("A7", "개별 몫 합계(주식 + 대기 현금, 원)"),
    ]
    for cell, text in labels:
        r = cell[1:]
        ws.merge_cells(f"A{r}:C{r}")
        ws.merge_cells(f"D{r}:E{r}")
        w.put(cell, text, bold=True)
    w.input_cells(["D"], 3, 4, FMT_INT)
    w.input_cells(["D"], 5, 5, FMT_DEC)
    w.note("D3", "지수 몫, 개별 몫, 현금을 모두 합친 투자 계좌 전체 평가액. 손실 한도(2%)의 기준.")
    w.note("D4", "개별 종목 몫으로 남겨 둔 현금. 비중 40% 계산에 들어간다.")
    first, last = HOLD_FIRST, HOLD_LAST
    w.put("D6", f"=COUNT(J{first}:J{last})", bold=True)
    w.put("D7", f"=SUM(J{first}:J{last})+N(D4)", fmt=FMT_INT, bold=True)
    rules = [
        "매도 판단 우선순위(먼저 걸린 것 하나를 제안)",
        "① 손실 한도: 손실액(원화 원가 − 원화 평가액) ≥ 계좌 평가액의 2%, 또는 현지 통화 가격 ≤ 매수 평균가 −25% → 매도",
        "② 핵심 매수 이유가 무너짐(매수이유 시트) → 매도 검토",
        "③ 3종목 이상 보유 중 한 종목이 개별 몫의 40% 초과 → 초과분만 매도",
        "④ 점수 50 미만이고 같은 국가에 10점 이상 높은 후보 → 교체 제안",
    ]
    for i, t in enumerate(rules, start=3):
        w.put(f"G{i}", t, bold=(i == 3))
    ws.merge_cells("A9:H9")
    ws.merge_cells("I9:W9")
    w.put("A9", "입력 (노란 칸)", bold=True, fill=F_GROUP)
    w.put("I9", "결과", bold=True, fill=F_KEY)
    w.header(
        10,
        [
            "국가", "종목코드", "종목명", "수량", "매수 평균가\n(현지 통화)", "원화 매수원가\n합계(원)", "현재가\n(현지 통화)", "최근 실적\n발표일",
            "최종 제안", "원화 평가액", "손실액(원)", "손실액\n/ 계좌", "매수가 대비", "① 손실 한도", "매수 이유 기록", "② 핵심 이유",
            "실적 점검", "비중\n(개별 몫)", "③ 비중 초과", "초과분(원)", "점수", "같은 국가\n최고 후보", "④ 교체",
        ],
    )
    ws.row_dimensions[10].height = 42
    a = rng(SH_REASON, "A", REASON_FIRST, REASON_LAST)
    for r in range(first, last + 1):
        best = f'IF(A{r}="한국",{ref(SH_STOCK["한국"], STOCK_BEST_SCORE)},{ref(SH_STOCK["미국"], STOCK_BEST_SCORE)})'
        best_name = f'IF(A{r}="한국",{ref(SH_STOCK["한국"], STOCK_BEST_NAME)},{ref(SH_STOCK["미국"], STOCK_BEST_NAME)})'
        fx_ok = f'OR(A{r}<>"미국",ISNUMBER({HOLD_FX}))'
        ws[f"J{r}"] = f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(G{r}),{fx_ok}),D{r}*G{r}*IF(A{r}="미국",{HOLD_FX},1),"")'
        ws[f"K{r}"] = f'=IF(AND(ISNUMBER(F{r}),ISNUMBER(J{r})),F{r}-J{r},"")'
        ws[f"L{r}"] = f'=IF(AND(ISNUMBER(K{r}),ISNUMBER({HOLD_ACC})),IF({HOLD_ACC}>0,K{r}/{HOLD_ACC},""),"")'
        ws[f"M{r}"] = f'=IF(AND(ISNUMBER(G{r}),ISNUMBER(E{r})),IF(E{r}>0,G{r}/E{r}-1,""),"")'
        ws[f"N{r}"] = (
            f'=IF(B{r}="","",IF(OR(AND(ISNUMBER(L{r}),L{r}>={CFG["loss_acc"]}),'
            f'AND(ISNUMBER(M{r}),M{r}<={CFG["loss_price"]})),"도달",""))'
        )
        ws[f"O{r}"] = f'=IF(B{r}="","",{record_status_expr(f"B{r}")})'
        ws[f"P{r}"] = f'=IF(B{r}="","",IF(SUMIFS({rng(SH_REASON, "M", REASON_FIRST, REASON_LAST)},{a},B{r})>0,"무너짐",""))'
        ws[f"Q{r}"] = (
            f'=IF(OR(B{r}="",H{r}=""),"",IF(COUNTIF({a},B{r})=0,"이유 기록 없음",'
            f'IF(SUMIFS({rng(SH_REASON, "O", REASON_FIRST, REASON_LAST)},{a},B{r})>0,"점검 필요","점검 완료")))'
        )
        ws[f"R{r}"] = f'=IF(AND(ISNUMBER(J{r}),{HOLD_SUM}>0),J{r}/{HOLD_SUM},"")'
        ws[f"S{r}"] = f'=IF(AND({HOLD_N}>={CFG["conc_n"]},ISNUMBER(R{r})),IF(R{r}>{CFG["conc"]},"초과",""),"")'
        ws[f"T{r}"] = f'=IF(S{r}="초과",J{r}-{CFG["conc"]}*{HOLD_SUM},"")'
        ws[f"U{r}"] = f"={score_lookup(f'A{r}', f'B{r}')}"
        ws[f"V{r}"] = f'=IF(B{r}="","",IF(ISNUMBER({best}),{best_name}&" ("&TEXT({best},"0.0")&")",""))'
        ws[f"W{r}"] = (
            f'=IF(AND(ISNUMBER(U{r}),ISNUMBER({best})),IF(AND(U{r}<{CFG["repl_below"]},'
            f'ROUND({best}-U{r},1)>={CFG["repl_gap"]}),"제안",""),"")'
        )
        ws[f"I{r}"] = (
            f'=IF(B{r}="","",IF(N{r}="도달","① 손실 한도 도달 → 매도",IF(P{r}="무너짐","② 핵심 매수 이유 붕괴 → 매도 검토",'
            f'IF(S{r}="초과","③ 비중 초과 → 초과분 "&TEXT(T{r},"#,##0")&"원 매도",'
            f'IF(W{r}="제안","④ 교체 제안 → "&{best_name},"유지")))))'
        )
        ws[f"I{r}"].font = Font(name=FONT, bold=True, size=10)
    w.input_cells(list("ABC"), first, last)
    w.fmt(["B"], first, last, "@")
    w.input_cells(["D"], first, last, FMT_INT)
    w.input_cells(["E", "G"], first, last, FMT_DEC)
    w.input_cells(["F"], first, last, FMT_INT)
    w.input_cells(["H"], first, last, FMT_DATE)
    w.fmt(["J", "K", "T"], first, last, FMT_INT)
    w.fmt(["L", "M", "R"], first, last, FMT_PCT)
    w.fmt(["U"], first, last, FMT_SCORE)
    w.validate(f"A{first}:A{last}", ["한국", "미국"])
    target = f"I{first}:I{last}"
    for mark, fill in (("①", F_RED), ("②", F_ORANGE), ("③", F_YELLOW), ("④", F_BLUE)):
        ws.conditional_formatting.add(target, FormulaRule(formula=[f'LEFT(I{first},1)="{mark}"'], fill=fill))
    w.widths(
        {
            "A": 6, "B": 10, "C": 16, "D": 8, "E": 11, "F": 14, "G": 11, "H": 11, "I": 36, "J": 13, "K": 12,
            "L": 8, "M": 9, "N": 9, "O": 16, "P": 9, "Q": 10, "R": 8, "S": 9, "T": 12, "U": 7, "V": 20, "W": 7,
        }
    )
    ws.freeze_panes = f"D{first}"


# ---------------------------------------------------------------- 매수 계획


def build_plan(ws):
    w = W(ws)
    w.title(
        "신규 종목 매수 계획 (명세 3·7)",
        "시장 심리 점수로 몇 번에 나눠 살지만 정한다(금액은 똑같이 나눔). 미국 종목은 미국 지수, 한국 종목은 한국 지수를 쓴다.",
    )
    w.header(
        5,
        [
            "국가", "종목코드", "종목명", "총 매수\n예정액(원)", "시작일", "제안", "시장 심리\n점수", "구간", "횟수", "간격(주)",
            "1회 금액(원)", "1회차", "2회차", "3회차", "4회차", "종목 점수", "편입 후보", "매수 이유 기록",
        ],
    )
    ws.row_dimensions[5].height = 32
    for r in range(PLAN_FIRST, PLAN_LAST + 1):
        ws[f"G{r}"] = f'=IF(A{r}="미국",{ref(SH_SENT["미국"], "$B$6")},IF(A{r}="한국",{ref(SH_SENT["한국"], "$B$6")},""))'
        ws[f"H{r}"] = f'=IF(ISNUMBER(G{r}),{band_lookup(f"G{r}", "name")},"")'
        ws[f"I{r}"] = f'=IF(ISNUMBER(G{r}),{band_lookup(f"G{r}", "n")},"")'
        ws[f"J{r}"] = f'=IF(ISNUMBER(G{r}),{band_lookup(f"G{r}", "w")},"")'
        # 금액과 날짜는 매수 이유 기록이 끝난 종목에만 보여 준다
        ws[f"K{r}"] = f'=IF(AND(R{r}="기록 완료",ISNUMBER(D{r}),ISNUMBER(I{r})),D{r}/I{r},"")'
        for k, c in enumerate("LMNO", start=1):
            ws[f"{c}{r}"] = f'=IF(AND(R{r}="기록 완료",ISNUMBER(E{r}),ISNUMBER(I{r})),IF({k}<=I{r},E{r}+{k - 1}*J{r}*7,""),"")'
        ws[f"P{r}"] = f"={score_lookup(f'A{r}', f'B{r}')}"
        ws[f"Q{r}"] = f'=IF(B{r}="","",IF(ISNUMBER(P{r}),IF(P{r}>={CFG["cand"]},"예","아니오(70 미만)"),"점수 없음"))'
        ws[f"R{r}"] = f'=IF(B{r}="","",{record_status_expr(f"B{r}")})'
        won = f'TEXT(K{r},"#,##0")'
        ws[f"F{r}"] = (
            f'=IF(B{r}="","",IF(R{r}<>"기록 완료","매수하지 않음: 매수 이유 "&R{r},IF(NOT(ISNUMBER(G{r})),"심리 점수 없음",'
            f'IF(I{r}=1,"1회에 전부"&IF(ISNUMBER(K{r}),", "&{won}&"원",""),'
            f'I{r}&"회, "&J{r}&"주 간격"&IF(ISNUMBER(K{r}),", 1회 "&{won}&"원",""))'
            f'&IF(Q{r}<>"예"," · 주의: 편입 후보 아님",""))))'
        )
        ws[f"F{r}"].font = Font(name=FONT, bold=True, size=10)
    w.input_cells(list("ABC"), PLAN_FIRST, PLAN_LAST)
    w.fmt(["B"], PLAN_FIRST, PLAN_LAST, "@")
    w.input_cells(["D"], PLAN_FIRST, PLAN_LAST, FMT_INT)
    w.input_cells(["E"], PLAN_FIRST, PLAN_LAST, FMT_DATE)
    w.fmt(["G", "P"], PLAN_FIRST, PLAN_LAST, FMT_SCORE)
    w.fmt(["K"], PLAN_FIRST, PLAN_LAST, FMT_INT)
    w.fmt(list("LMNO"), PLAN_FIRST, PLAN_LAST, FMT_DATE)
    w.validate(f"A{PLAN_FIRST}:A{PLAN_LAST}", ["한국", "미국"])
    w.widths({"A": 6, "B": 10, "C": 16, "D": 13, "E": 11, "F": 40, "G": 9, "H": 9, "I": 6, "J": 7, "K": 12,
              "L": 11, "M": 11, "N": 11, "O": 11, "P": 8, "Q": 14, "R": 16})
    ws.freeze_panes = f"D{PLAN_FIRST}"


# ---------------------------------------------------------------- 판단 기록


def build_log(ws):
    w = W(ws)
    w.title(
        "판단 기록",
        "제안을 보고 결정할 때마다 한 줄. 대시보드의 제안을 복사해 '값만 붙여넣기' 하고, 내 결정과 이유를 적는다. 나중에 판단을 돌아보는 용도.",
    )
    w.header(4, ["날짜", "국가", "종목코드", "종목명", "프로그램 제안", "점수", "시장 심리", "내 결정", "실제 행동", "금액(원)", "이유·메모"])
    w.input_cells(list("ABCDEFGHIJK"), 5, 504)
    w.fmt(["A"], 5, 504, FMT_DATE)
    w.fmt(["C"], 5, 504, "@")
    w.fmt(["F", "G"], 5, 504, FMT_SCORE)
    w.fmt(["J"], 5, 504, FMT_INT)
    w.validate("B5:B504", ["한국", "미국"])
    w.validate("H5:H504", ["따름", "따르지 않음", "보류"])
    w.validate("I5:I504", ["매수", "매도", "일부 매도", "교체", "유지"])
    w.widths({"A": 11, "B": 6, "C": 10, "D": 16, "E": 36, "F": 7, "G": 9, "H": 11, "I": 10, "J": 13, "K": 50})
    ws.freeze_panes = "A5"


# ---------------------------------------------------------------- 시장 심리

# 원자료 블록: (키, 제목, 날짜 열, 값 열, 구글 시트 자동 입력용 FRED id 또는 None)
RAW_BLOCKS = {
    "미국": (
        ("sp500", "S&P 500 지수 — FRED SP500", "A", "B", "SP500"),
        ("vix", "VIX — FRED VIXCLS", "D", "E", "VIXCLS"),
        ("baa", "Baa 회사채 − 미 10년물 금리차(%p) — FRED BAA10Y", "G", "H", "BAA10Y"),
    ),
    "한국": (
        ("kospi", "코스피 지수 — 한국은행 ECOS 802Y001 또는 KRX", "A", "B", None),
        ("vkospi", "VKOSPI — KRX 정보데이터시스템", "D", "E", None),
        ("aa", "회사채 3년 AA- 금리(%) — ECOS 817Y002", "G", "H", None),
        ("ktb", "국고채 3년 금리(%) — ECOS 817Y002", "J", "K", None),
        ("credit", "신용융자잔고 — 금융투자협회 FreeSIS", "M", "N", None),
    ),
}

RAW_HOWTO = {
    "미국": [
        "붙여넣는 법: 각 블록의 2행(머리글 줄)부터 '날짜, 값' 두 열을 붙여 넣는다. 순서는 상관없다.",
        "기간: 최근 6년 이상(5년 + 125거래일). 행이 4,000을 넘지 않게 한다.",
        "빈 값('.' 또는 빈칸)은 그대로 두어도 된다(그날은 직전 값을 씀).",
        "구글 시트 자동 입력(선택): 각 블록 2행 첫 칸(A2, D2, G2)에 아래 수식을 넣으면 FRED에서 바로 가져온다.",
        '=IMPORTDATA("https://fred.stlouisfed.org/graph/fredgraph.csv?id=SP500&cosd="&TEXT(EDATE(TODAY(),-72),"yyyy-mm-dd"))',
        "  VIX는 id=VIXCLS, 금리차는 id=BAA10Y 로 바꾼다. (엑셀에서는 동작하지 않는다)",
    ],
    "한국": [
        "붙여넣는 법: 각 블록의 2행(머리글 줄)부터 '날짜, 값' 두 열을 붙여 넣는다. 최신순이어도 된다.",
        "기간: 최근 6년 이상. 행이 4,000을 넘지 않게 한다. 날짜는 날짜 형식이어야 한다(글자 날짜면 계산에서 빠짐).",
        "코스피: ECOS 통계표 802Y001 주식시장(일) → KOSPI지수, 또는 KRX [지수 → 주가지수 → 개별지수 시세 추이].",
        "VKOSPI: KRX 정보데이터시스템 [지수 → 파생상품지수 → 개별지수 시세 추이] → 코스피 200 변동성지수.",
        "금리: ECOS 817Y002 시장금리(일별) → 회사채(3년, AA-), 국고채(3년). 금리차 = AA- − 국고채.",
        "신용융자잔고: 금융투자협회 FreeSIS [주식 → 신용공여 → 신용공여 잔고 추이] → 신용거래융자 합계(유가+코스닥).",
    ],
}

# 달력 열: (키, 머리글, 종류, 원자료 키, 서식)
CAL_COLS = {
    "미국": (
        ("date", "날짜", "date", None, FMT_DATE),
        ("base", "S&P 500", "base", "sp500", FMT_DEC),
        ("ma", "125일 평균 대비", "ma", None, FMT_PCT),
        ("vix", "VIX", "ffill", "vix", FMT_DEC),
        ("baa", "Baa − 10년물", "ffill", "baa", FMT_DEC),
    ),
    "한국": (
        ("date", "날짜", "date", None, FMT_DATE),
        ("base", "코스피", "base", "kospi", FMT_DEC),
        ("ma", "125일 평균 대비", "ma", None, FMT_PCT),
        ("vkospi", "VKOSPI", "ffill", "vkospi", FMT_DEC),
        ("aa", "회사채 AA- 3년", "ffill", "aa", "0.000"),
        ("ktb", "국고채 3년", "ffill", "ktb", "0.000"),
        ("spread", "금리차", "spread", None, "0.000"),
        ("credit_level", "신용융자잔고", "ffill", "credit", FMT_INT),
        ("credit", "20일 증감률", "credit", None, "0.00%"),
    ),
}

# 심리 지표: (달력 키, 이름, 공포 방향 높을수록?, 공포 방향 설명, 최종 자료일 원자료 키들)
SENT_IND = {
    "미국": (
        ("vix", "VIX", True, "높을수록 공포", ("vix",)),
        ("baa", "Baa − 미 10년물", True, "넓을수록 공포", ("baa",)),
        ("ma", "S&P 500 / 125일 평균", False, "낮을수록 공포", ("sp500",)),
    ),
    "한국": (
        ("vkospi", "VKOSPI", True, "높을수록 공포", ("vkospi",)),
        ("spread", "회사채 AA- − 국고채 3년", True, "넓을수록 공포", ("aa", "ktb")),
        ("ma", "코스피 / 125일 평균", False, "낮을수록 공포", ("kospi",)),
        ("credit", "신용융자잔고 20일 증감률", False, "줄어들수록 공포", ("credit",)),
    ),
}


def build_raw(ws, title, blocks, howto, rows_for):
    w = W(ws)
    last_col = None
    for key, label, dc, vc, _ in blocks:
        w.put(f"{dc}1", label, bold=True)
        w.header(2, ["날짜", "값"], start=ci(dc))
        n = rows_for(key)
        w.input_cells([dc], RAW_FIRST, RAW_FIRST + n - 1, FMT_DATE)
        w.input_cells([vc], RAW_FIRST, RAW_FIRST + n - 1)
        w.widths({dc: 12, vc: 12})
        last_col = vc
    note_col = col(ci(last_col) + 2)
    w.put(f"{note_col}1", title, bold=True, size=12)
    for i, t in enumerate(howto, start=2):
        c = w.put(f"{note_col}{i}", t, italic=not t.startswith("="), color="595959")
        c.data_type = "s"  # '='로 시작하는 안내 문구가 수식으로 저장되지 않게
    ws.column_dimensions[note_col].width = 90
    ws.freeze_panes = "A3"


def build_sentiment(ws, market: str):
    w = W(ws)
    raw = SH_RAW[market]
    blocks = {b[0]: b for b in RAW_BLOCKS[market]}
    cols = CAL_COLS[market]
    inds = SENT_IND[market]
    w.title(
        f"시장 심리 지수 — {market} (명세 3)",
        "신규 종목 매수를 몇 번에 나눌지만 정한다. 매도와 지수 몫에는 쓰지 않는다. 0 = 극단 공포, 100 = 극단 탐욕.",
    )

    # 열 배치
    letter = {}
    i = 1
    for key, *_ in cols:
        letter[key] = col(i)
        i += 1
    letter["inw"] = col(i)
    i += 1
    for key, *_ in inds:
        letter["w_" + key] = col(i)
        i += 1
    for key, *_ in inds:
        letter["g_" + key] = col(i)
        i += 1
    letter["comp"] = col(i)
    i += 2
    helper = {}
    for key in blocks:
        helper[key] = col(i)
        i += 1
    help_first = SENT_FIRST
    help_last = SENT_FIRST + RAW_DAILY - 1

    def H(key):
        return f"${helper[key]}${help_first}:${helper[key]}${help_last}"

    def RAWV(key):
        return rng(raw, blocks[key][3], RAW_FIRST, RAW_FIRST + RAW_DAILY - 1)

    def CALR(key):
        return f"${letter[key]}${SENT_FIRST}:${letter[key]}${SENT_LAST}"

    idx_key = cols[1][3]

    # 보조열: 원자료 i행 → 날짜(값이 숫자일 때만)
    for key, _, dc, vc, _ in RAW_BLOCKS[market]:
        for rr in range(RAW_FIRST, RAW_FIRST + RAW_DAILY):
            h = rr + SENT_HELP_OFFSET
            d, v = ref(raw, f"${dc}{rr}"), ref(raw, f"${vc}{rr}")
            ws[f"{helper[key]}{h}"] = f'=IF(AND(ISNUMBER({d}),ISNUMBER({v})),{d},"")'

    # 요약
    w.put("A4", "기준일(지수 마지막 거래일)", bold=True)
    w.put("A5", "5년 창 시작(이 날 다음부터)", bold=True)
    w.put("A6", "최종 점수", bold=True)
    w.put("A7", "구간", bold=True)
    w.put("A8", "신규 매수", bold=True)
    w.put("A9", "오늘 지표 평균", bold=True)
    w.put("A10", "점검", bold=True)
    w.put("D4", "지수 원자료 행 수")
    w.put("D5", "달력 행 수")
    w.put("D6", "건너뛴 오래된 행")
    w.put("D7", "창 안 거래일 수")
    w.put("D8", "지표 평균이 있는 날")
    w.put("D9", "원자료 최대 행 수")
    w.put("E4", f"=COUNT({H(idx_key)})")
    w.put("E5", f"=MIN(E4,{CAL_ROWS})")
    w.put("E6", f"=MAX(E4-{CAL_ROWS},0)")
    w.put("E7", f"=SUM({CALR('inw')})")
    w.put("E8", f"=COUNT({CALR('comp')})")
    counts = ",".join(f"COUNTA({rng(raw, b[2], RAW_FIRST, RAW_FIRST + RAW_DAILY - 1)})" for b in RAW_BLOCKS[market])
    w.put("E9", f"=MAX({counts})")
    w.put("B4", f'=IF(E5=0,"",MAX({CALR("date")}))', fmt=FMT_DATE, bold=True)
    w.put("B5", f'=IF(B4="","",EDATE(B4,-{CFG["sent_m"]}))', fmt=FMT_DATE)
    w.put("B9", f'=IF(B4="","",INDEX({CALR("comp")},$E$5))', fmt=FMT_SCORE)
    w.put("B6", f'=IF(ISNUMBER(B9),ROUND({pct_expr("B9", CALR("comp"), "$E$8")},1),"")', fmt=FMT_SCORE, bold=True, fill=F_KEY, size=14)
    w.put("B7", f'=IF(ISNUMBER(B6),{band_lookup("B6", "name")},"")', bold=True, fill=F_KEY)
    times, weeks = band_lookup("B6", "n"), band_lookup("B6", "w")
    w.put("B8", f'=IF(ISNUMBER(B6),IF({times}=1,"1회에 전부",{times}&"회, "&{weeks}&"주 간격"),"")', bold=True, fill=F_KEY)

    tr = 12
    w.header(tr, ["지표", "최신값", "최종 자료일", "5년 백분위\n(값 기준)", "공포 방향", "탐욕 쪽 점수\n(반전 후)", "창 안\n자료 수", "점검"])
    ws.row_dimensions[tr].height = 32
    for j, (key, label, fear_high, fear_text, sources) in enumerate(inds):
        r = tr + 1 + j
        w.put(f"A{r}", label, bold=True)
        w.put(f"B{r}", f'=IF($B$4="","",INDEX({CALR(key)},$E$5))', fmt=dict((c[0], c[4]) for c in cols)[key])
        if len(sources) == 1:
            s = sources[0]
            last_src = f'IF(COUNT({H(s)})=0,"",MAX({H(s)}))'
        else:
            cond = ",".join(f"COUNT({H(s)})=0" for s in sources)
            last_src = f'IF(OR({cond}),"",MIN({",".join(f"MAX({H(s)})" for s in sources)}))'
        w.put(f"C{r}", f"={last_src}", fmt=FMT_DATE)
        w.put(f"G{r}", f"=COUNT({CALR('w_' + key)})")
        w.put(f"D{r}", f'=IF(AND(ISNUMBER(B{r}),G{r}>0),{pct_expr(f"B{r}", CALR("w_" + key), f"G{r}")},"")', fmt="0")
        w.put(f"E{r}", fear_text)
        w.put(f"F{r}", f'=IF($B$4="","",INDEX({CALR("g_" + key)},$E$5))', fmt="0")
        w.put(
            f"H{r}",
            f'=IF($B$4="","",_xlfn.TEXTJOIN(", ",TRUE,IF(G{r}<$E$7,"{label} 5년 자료 부족",""),'
            f'IF(OR(C{r}="",C{r}<$B$4-{CFG["stale"]}),"{label} 자료 지연","")))',
        )
    last_tr = tr + len(inds)
    joined = f'_xlfn.TEXTJOIN(", ",TRUE,H{tr + 1}:H{last_tr},IF(E9>={RAW_DAILY},"원자료 {RAW_DAILY:,}행 꽉 참",""))'
    w.put("B10", f'=IF(B4="","원자료 없음",IF({joined}="","이상 없음",{joined}))')

    # 달력 표
    headers = [c[1] for c in cols] + ["창 안"] + [f"{lab}\n(창 안 값)" for _, lab, *_ in inds] + [f"{lab}\n탐욕 점수" for _, lab, *_ in inds] + ["지표 평균"]
    w.header(SENT_HDR, headers)
    ws.row_dimensions[SENT_HDR].height = 42
    w.header(SENT_HDR, [f"보조: {b[0]} 날짜" for b in RAW_BLOCKS[market]], start=ci(helper[RAW_BLOCKS[market][0][0]]))
    ma_n, lag = SPEC["ma_days"], SPEC["credit_lag"]
    for k in range(1, CAL_ROWS + 1):
        r = SENT_FIRST + k - 1
        A = f"{letter['date']}{r}"
        for key, _, kind, src, _ in cols:
            cell = f"{letter[key]}{r}"
            if kind == "date":
                ws[cell] = f'=IF({k}>$E$5,"",SMALL({H(idx_key)},$E$6+{k}))'
            elif kind == "base":
                ws[cell] = f'=IF({A}="","",INDEX({RAWV(src)},MATCH({A},{H(src)},0)))'
            elif kind == "ma":
                # 앞쪽 124행도 빈 칸이 아니라 "" 수식으로 둔다(INDEX는 빈 칸을 0으로 돌려준다)
                b = letter["base"]
                ws[cell] = f'=IF({A}="","",{b}{r}/AVERAGE({b}{r - ma_n + 1}:{b}{r})-1)' if k >= ma_n else '=""'
            elif kind == "ffill":
                prev = f"{letter[key]}{r - 1}" if k > 1 else '""'
                ws[cell] = f'=IF({A}="","",IFERROR(INDEX({RAWV(src)},MATCH({A},{H(src)},0)),{prev}))'
            elif kind == "spread":
                aa, kt = f"{letter['aa']}{r}", f"{letter['ktb']}{r}"
                ws[cell] = f'=IF(AND(ISNUMBER({aa}),ISNUMBER({kt})),ROUND({aa}-{kt},{IMPL["tie_digits"]}),"")'
            elif kind == "credit":
                now, past = f"{letter['credit_level']}{r}", f"{letter['credit_level']}{r - lag}"
                ws[cell] = f'=IF(AND(ISNUMBER({now}),ISNUMBER({past})),IF({past}>0,{now}/{past}-1,""),"")' if k > lag else '=""'
            ws[cell].number_format = dict((c[0], c[4]) for c in cols)[key]
        ws[f"{letter['inw']}{r}"] = f'=IF({A}="",0,IF({A}>$B$5,1,0))'
        for j, (key, label, fear_high, *_) in enumerate(inds):
            wc, gc = letter["w_" + key], letter["g_" + key]
            val = f"{letter[key]}{r}"
            ws[f"{wc}{r}"] = f'=IF(AND({letter["inw"]}{r}=1,ISNUMBER({val})),{val},"")'
            p = pct_expr(f"{wc}{r}", CALR("w_" + key), f"$G${tr + 1 + j}")
            ws[f"{gc}{r}"] = f'=IF({wc}{r}="","",{"100-" if fear_high else ""}{p})'
            ws[f"{gc}{r}"].number_format = "0.0"
            ws[f"{wc}{r}"].number_format = dict((c[0], c[4]) for c in cols)[key]
        g1, g2 = letter["g_" + inds[0][0]], letter["g_" + inds[-1][0]]
        ws[f"{letter['comp']}{r}"] = f'=IF(COUNT({g1}{r}:{g2}{r})={len(inds)},ROUND(AVERAGE({g1}{r}:{g2}{r}),{IMPL["tie_digits"]}),"")'
        ws[f"{letter['comp']}{r}"].number_format = "0.0"
    widths = {letter[k]: 11 for k in letter}
    widths.update({"A": 26, "B": 12, "C": 12, "D": 12, "E": 14, "F": 12, "G": 9, "H": 30})
    w.widths(widths)
    w.hide(ci(helper[RAW_BLOCKS[market][0][0]]), ci(helper[RAW_BLOCKS[market][-1][0]]))


# ---------------------------------------------------------------- 거시

MACRO_RAW = {
    "sahm": ("A", "B", RAW_MONTHLY, "SAHMREALTIME"),
    "dgs10": ("D", "E", RAW_DAILY, "DGS10"),
    "t10y3m": ("G", "H", RAW_DAILY, "T10Y3M"),
    "jpy": ("J", "K", RAW_DAILY, "DEXJPUS"),
    "exports": ("M", "N", RAW_MONTHLY, "XTEXVA01KRM667N"),
}


def build_macro(ws):
    w = W(ws)
    w.title("거시 참고 패널 (명세 4)", "표시만 한다. 어떤 규칙의 입력도 아니고 합산 점수도 없다.")
    w.header(4, ["지표", "보는 것", "현재값", "단위", "기준일", "10년 백분위", "창 안\n자료 수", "창 첫 자료일", "출처", "점검", "원자료 날짜 수", "창 시작", "정렬 오류", "점검 문구"])
    ws.row_dimensions[4].height = 32
    raw = SH_MACRO_RAW
    i = ci("P")
    helper_cols = {}
    for key, *_ in MACRO_SERIES:
        kind = next(s[3] for s in MACRO_SERIES if s[0] == key)
        names = {"level": ("x", "xd", "w"), "diff": ("cv", "x", "xd", "w"), "ratio": ("cv", "x", "xd", "w"), "yoy3": ("n", "x", "xd", "w")}[kind]
        helper_cols[key] = {}
        for nm in names:
            helper_cols[key][nm] = col(i)
            i += 1
    for row, (key, label, what, kind, unit, source) in enumerate(MACRO_SERIES, start=5):
        dc, vc, nrows, fred = MACRO_RAW[key]
        h = helper_cols[key]
        hf, hl = MACRO_HELP_FIRST, MACRO_HELP_FIRST + nrows - 1
        rf, rl = RAW_FIRST, RAW_FIRST + nrows - 1
        rd_all = rng(raw, dc, rf, rl)
        rv_all = rng(raw, vc, rf, rl)

        def HC(nm):
            return f"${h[nm]}${hf}:${h[nm]}${hl}"

        w.header(MACRO_HELP_FIRST - 1, [f"{key}:{nm}" for nm in h], start=ci(h[next(iter(h))]))
        for rr in range(rf, rl + 1):
            hr = rr + MACRO_HELP_OFFSET
            rd, rv = ref(raw, f"${dc}{rr}"), ref(raw, f"${vc}{rr}")
            ok = f"AND(ISNUMBER({rd}),ISNUMBER({rv}))"
            if kind == "level":
                ws[f"{h['x']}{hr}"] = f'=IF({ok},{rv},"")'
            elif kind in ("diff", "ratio"):
                prev = f"{h['cv']}{hr - 1}" if rr > rf else '""'
                ws[f"{h['cv']}{hr}"] = f"=IF(ISNUMBER({rv}),{rv},{prev})"
                dyn = f"{ref(raw, f'${dc}${rf}')}:INDEX({rd_all},$K${row})"
                lagv = f"INDEX({HC('cv')},MATCH(EDATE({rd},-{CFG['chg_m']}),{dyn},1))"
                expr = f"{rv}-{lagv}" if kind == "diff" else f"{rv}/{lagv}-1"
                ws[f"{h['x']}{hr}"] = f'=IF({ok},IFERROR({expr},""),"")'
            else:  # yoy3
                ws[f"{h['n']}{hr}"] = f"=IF({ok},1,0)"
                m = CFG["chg_m"]
                win1 = f'{rd_all},">"&EDATE({rd},-{m}),{rd_all},"<="&{rd}'
                win0 = f'{rd_all},">"&EDATE({rd},-12-{m}),{rd_all},"<="&EDATE({rd},-12)'
                ws[f"{h['x']}{hr}"] = (
                    f'=IF({ok},IF(AND(SUMIFS({HC("n")},{win1})=3,SUMIFS({HC("n")},{win0})=3),'
                    f'IFERROR(SUMIFS({rv_all},{win1})/SUMIFS({rv_all},{win0})-1,""),""),"")'
                )
            ws[f"{h['xd']}{hr}"] = f'=IF(ISNUMBER({h["x"]}{hr}),{rd},"")'
            ws[f"{h['w']}{hr}"] = f'=IF(ISNUMBER({h["xd"]}{hr}),IF({h["xd"]}{hr}>$L${row},{h["x"]}{hr},""),"")'
        val_fmt = "0.00" if unit == "%p" else FMT_PCT
        w.put(f"A{row}", label, bold=True)
        w.put(f"B{row}", what)
        w.put(f"E{row}", f'=IF(COUNT({HC("xd")})=0,"",MAX({HC("xd")}))', fmt=FMT_DATE)
        w.put(f"C{row}", f'=IF(E{row}="","",INDEX({HC("x")},MATCH(E{row},{HC("xd")},0)))', fmt=val_fmt, bold=True)
        w.put(f"D{row}", unit)
        w.put(f"L{row}", f'=IF(E{row}="","",EDATE(E{row},-{CFG["macro_m"]}))', fmt=FMT_DATE)
        w.put(f"G{row}", f"=COUNT({HC('w')})")
        w.put(f"F{row}", f'=IF(G{row}=0,"",{pct_expr(f"C{row}", HC("w"), f"G{row}")})', fmt="0", bold=True)
        w.put(f"H{row}", f'=IF(G{row}=0,"",_xlfn.MINIFS({HC("xd")},{HC("xd")},">"&L{row}))', fmt=FMT_DATE)
        w.put(f"I{row}", source)
        w.put(f"K{row}", f"=COUNT({rd_all})")
        a_rng = rng(raw, dc, rf, rl - 1)
        b_rng = rng(raw, dc, rf + 1, rl)
        w.put(f"M{row}", f"=SUMPRODUCT(({a_rng}>={b_rng})*ISNUMBER({b_rng}))")
        w.put(
            f"N{row}",
            f'=_xlfn.TEXTJOIN(", ",TRUE,IF(E{row}="","자료 없음",""),IF(M{row}>0,"날짜 정렬 오류",""),'
            f'IF(G{row}>0,IF(H{row}>L{row}+{CFG["slack"]},"10년 자료 부족",""),""))',
        )
        w.put(f"J{row}", f'=IF(N{row}="","이상 없음",N{row})')
    w.put("A11", "· 현재값은 가장 최근 자료, 10년 백분위는 최근 10년 값 가운데 위치(0 = 가장 낮음, 100 = 가장 높음).", italic=True, color="595959")
    w.put("A12", "· 미 10년물 3개월 변화 = 오늘 − 3개월 전(%p). 엔/달러 변화율이 음수면 엔화 강세. 한국 수출 = 최근 3개월 합 ÷ 1년 전 같은 3개월 합 − 1.", italic=True, color="595959")
    w.widths({"A": 26, "B": 18, "C": 10, "D": 6, "E": 11, "F": 10, "G": 8, "H": 12, "I": 24, "J": 22})
    w.hide(ci("K"), ci("N"))
    w.hide(ci("P"), i - 1)


MACRO_HOWTO = [
    "붙여넣는 법: 각 블록의 2행(머리글 줄)부터 FRED에서 받은 '날짜, 값' 두 열을 붙여 넣는다.",
    "날짜는 오름차순(FRED 기본)이어야 한다. 빈 값('.' 또는 빈칸)은 그대로 둔다.",
    "기간: 최근 11년 이상(10년 + 변화 계산용 여유). 일별 4,000행, 월별 600행까지.",
    "구글 시트 자동 입력(선택): 각 블록 2행 첫 칸(A2, D2, G2, J2, M2)에 아래 수식을 넣는다. id만 바꾼다.",
    '=IMPORTDATA("https://fred.stlouisfed.org/graph/fredgraph.csv?id=SAHMREALTIME&cosd="&TEXT(EDATE(TODAY(),-138),"yyyy-mm-dd"))',
    "  id: SAHMREALTIME, DGS10, T10Y3M, DEXJPUS, XTEXVA01KRM667N (엑셀에서는 동작하지 않는다)",
]


# ---------------------------------------------------------------- 대시보드


DASH_HOLD, DASH_CAND, DASH_MACRO = 11, 24, 37  # 각 표의 첫 줄


def build_dashboard(ws):
    w = W(ws)
    w.title("대시보드", "프로그램은 제안만 한다. 매매는 사람이 승인하고 직접 한다. 자동 매매 없음.")
    w.put("A4", "시장 심리 — 신규 종목 매수 분할에만 사용", bold=True, size=12)
    w.header(5, ["시장", "점수", "구간", "신규 매수", "기준일", "점검"])
    for r, m in ((6, "미국"), (7, "한국")):
        sh = SH_SENT[m]
        w.put(f"A{r}", m, bold=True)
        w.put(f"B{r}", f"={ref(sh, '$B$6')}", fmt=FMT_SCORE, bold=True, fill=F_KEY)
        w.put(f"C{r}", f"={ref(sh, '$B$7')}", bold=True, fill=F_KEY)
        w.put(f"D{r}", f"={ref(sh, '$B$8')}", bold=True)
        w.put(f"E{r}", f"={ref(sh, '$B$4')}", fmt=FMT_DATE)
        w.put(f"F{r}", f"={ref(sh, '$B$10')}")

    n_hold = HOLD_LAST - HOLD_FIRST + 1
    w.put(f"A{DASH_HOLD - 2}", "보유 종목 제안 — 우선순위 ① 손실 한도 ② 핵심 이유 붕괴 ③ 비중 초과 ④ 교체", bold=True, size=12)
    w.header(DASH_HOLD - 1, ["국가", "종목코드", "종목명", "최종 제안", "점수", "매수 이유 기록", "실적 점검"])
    for k in range(n_hold):
        r, hr = DASH_HOLD + k, HOLD_FIRST + k
        for c, src in zip("ABCDEFG", "ABCIUOQ"):
            text = '&""' if src in "ABC" else ""  # 입력 칸이 비어 있으면 0이 아니라 빈 글자
            ws[f"{c}{r}"] = f'=IF({ref(SH_HOLD, f"$B{hr}")}="","",{ref(SH_HOLD, f"{src}{hr}")}{text})'
        ws[f"E{r}"].number_format = FMT_SCORE
        ws[f"D{r}"].font = Font(name=FONT, bold=True, size=10)
    target = f"D{DASH_HOLD}:D{DASH_HOLD + n_hold - 1}"
    for mark, fill in (("①", F_RED), ("②", F_ORANGE), ("③", F_YELLOW), ("④", F_BLUE)):
        ws.conditional_formatting.add(target, FormulaRule(formula=[f'LEFT(D{DASH_HOLD},1)="{mark}"'], fill=fill))

    w.put(f"A{DASH_CAND - 2}", "편입 후보 — 점수 70 이상, 미보유 (점수순)", bold=True, size=12)
    w.header(DASH_CAND - 1, ["순위", "점수", "한국 종목", "한 줄 요약", "점수", "미국 종목", "한 줄 요약"])
    for k in range(1, 11):
        r = DASH_CAND + k - 1
        w.put(f"A{r}", k)
        for c_score, c_name, c_sum, m in (("B", "C", "D", "한국"), ("E", "F", "G", "미국")):
            sh = SH_STOCK[m]
            key = rng(sh, SC["rankkey"], STOCK_FIRST, STOCK_LAST)
            pos = f"MATCH(LARGE({key},{k}),{key},0)"
            for c, field, text in ((c_name, "name", '&""'), (c_score, "score", ""), (c_sum, "summary", '&""')):
                ws[f"{c}{r}"] = f'=IFERROR(INDEX({rng(sh, SC[field], STOCK_FIRST, STOCK_LAST)},{pos}){text},"")'
            ws[f"{c_score}{r}"].number_format = FMT_SCORE

    mt = DASH_MACRO - 1
    w.put(f"A{mt - 1}", "거시 참고 패널 — 표시만, 어떤 규칙에도 쓰지 않음", bold=True, size=12)
    w.header(mt, ["지표", "", "", "보는 것", "현재값", "10년 백분위", "기준일 · 점검"])
    ws.merge_cells(f"A{mt}:C{mt}")
    for j, item in enumerate(MACRO_SERIES):
        r, sr = DASH_MACRO + j, 5 + j
        ws.merge_cells(f"A{r}:C{r}")
        ws[f"A{r}"] = f"={ref(SH_MACRO, f'A{sr}')}"
        ws[f"D{r}"] = f"={ref(SH_MACRO, f'B{sr}')}"
        ws[f"E{r}"] = f"={ref(SH_MACRO, f'C{sr}')}"
        ws[f"F{r}"] = f"={ref(SH_MACRO, f'F{sr}')}"
        e, j_ = ref(SH_MACRO, f"E{sr}"), ref(SH_MACRO, f"J{sr}")
        ws[f"G{r}"] = f'=IF({e}="",{j_},TEXT({e},"yyyy-mm-dd")&" · "&{j_})'
        ws[f"E{r}"].number_format = '0.00"%p"' if item[4] == "%p" else FMT_PCT
        ws[f"E{r}"].font = Font(name=FONT, bold=True, size=10)
        ws[f"F{r}"].number_format = "0"
    w.widths({"A": 8, "B": 11, "C": 18, "D": 38, "E": 11, "F": 18, "G": 38})


# ---------------------------------------------------------------- 안내

GUIDE = [
    ("투자 판단 보조 시트", "title"),
    ("판단 기준 명세(docs/spec.md)를 스프레드시트 수식으로 옮긴 1차 구현. 프로그램은 제안만 한다. 매매는 사람이 승인하고 직접 한다.", ""),
    ("", ""),
    ("칸 색", "head"),
    ("노란 바탕 + 파란 글씨 = 입력하는 칸. 나머지는 수식이므로 고치지 않는다. 숨긴 열은 중간 계산이다(열 숨기기 해제로 볼 수 있음).", ""),
    ("", ""),
    ("쓰는 순서", "head"),
    ("1. 종목_한국 / 종목_미국: 보유 종목과 후보의 재무·주가를 넣는다. → 거름망, 점수, 판정, 점수 설명 카드가 나온다.", ""),
    ("2. 매수이유: 보유 종목과 살 종목의 이유(1~3개, 핵심 1개), 근거, 무너짐 기준, 점검 결과를 적는다.", ""),
    ("3. 보유: 계좌 평가액, 대기 현금, 환율, 보유 종목을 넣는다. → 손실 한도·매도 판단을 우선순위대로 제안한다.", ""),
    ("4. 자료_미국심리 / 자료_한국심리: 지표 원자료를 붙여 넣는다. → 심리_미국 / 심리_한국에서 점수와 분할 매수 방법이 나온다.", ""),
    ("5. 매수계획: 새로 살 종목과 금액, 시작일을 넣는다. → 몇 회, 몇 주 간격, 1회 금액, 회차별 날짜가 나온다.", ""),
    ("6. 자료_거시: FRED 원자료를 붙여 넣는다. → 거시 탭에 현재값과 10년 백분위가 표시된다(규칙에는 쓰지 않음).", ""),
    ("7. 대시보드에서 한눈에 본다. 결정을 내리면 판단기록에 한 줄 남긴다.", ""),
    ("", ""),
    ("갱신 주기(권장)", "head"),
    ("재무(종목 시트 F~N열): 분기 실적 발표 뒤. 주가·시총(L, O, P열): 판단할 때마다 같은 기준일로.", ""),
    ("보유 현재가·환율·계좌 평가액: 월 1회, 그리고 매매 전.", ""),
    ("심리 원자료: 신규 종목을 사기 직전. 거시 원자료: 월 1회.", ""),
    ("매수 이유 점검: 실적 발표 때마다(보유 시트에 최근 실적 발표일을 넣으면 '점검 필요'를 알려 준다).", ""),
    ("", ""),
    ("예외 처리(구현하는 쪽이 정한 것)", "head"),
    ("· 거름망 항목 자료가 비어 있으면 통과로 보지 않는다(점수 없음).", ""),
    ("· 지표 4개 중 하나라도 계산할 수 없으면(자료 없음, 분모 0 이하) 점수를 매기지 않고 이유를 보여 준다.", ""),
    ("· 동점은 평균 순위. 백분위 = (순위 − 1) ÷ (종목 수 − 1) × 100, 종목이 하나면 50.", ""),
    ("· 교체 후보가 여럿 같은 점수면 시트 위쪽 종목을 보여 준다.", ""),
    ("· 심리 지표는 그날 값이 없으면 직전 값을 쓴다. 최종 자료일이 7일 넘게 오래되면 '자료 지연'으로 알린다.", ""),
    ("· 주가는 수정주가(분할·병합 반영)를 넣는다. 배당은 넣지 않는다(명세: 주가 수익률).", ""),
    ("", ""),
    ("하지 않는 것", "head"),
    ("자동 매매 · 거시 지표로 규칙 바꾸기 · 시장 타이밍으로 현금 비중 조절 · 이벤트 점수 · 명세에 없는 지표 추가", ""),
]


def build_guide(ws, example: bool):
    w = W(ws)
    r = 1
    for text, kind in GUIDE:
        if kind == "title":
            w.put(f"A{r}", text, bold=True, size=16)
        elif kind == "head":
            w.put(f"A{r}", text, bold=True, size=12)
        else:
            w.put(f"A{r}", text)
        r += 1
    if example:
        w.put(f"A{r + 1}", "이 파일은 가상 데이터로 채운 예시다. 종목·숫자·원자료는 모두 지어낸 값이다. 실제로 쓸 때는 빈 양식(투자판단.xlsx)을 쓴다.", bold=True, color="C00000")
    ws.column_dimensions["A"].width = 140


# ---------------------------------------------------------------- 조립


def build(path, data=None):
    """data가 있으면 입력 칸을 채운 예시 파일을 만든다."""
    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name=FONT, size=10)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    order = [
        SH_GUIDE, SH_DASH, SH_STOCK["한국"], SH_STOCK["미국"], SH_HOLD, SH_REASON, SH_PLAN, SH_LOG,
        SH_SENT["미국"], SH_SENT["한국"], SH_MACRO, SH_RAW["미국"], SH_RAW["한국"], SH_MACRO_RAW, SH_SET,
    ]
    sheets = {}
    for i, name in enumerate(order):
        sheets[name] = wb.active if i == 0 else wb.create_sheet()
        sheets[name].title = name

    build_guide(sheets[SH_GUIDE], example=data is not None)
    build_dashboard(sheets[SH_DASH])
    for country in ("한국", "미국"):
        build_stock(sheets[SH_STOCK[country]], country)
    build_holdings(sheets[SH_HOLD])
    build_reasons(sheets[SH_REASON])
    build_plan(sheets[SH_PLAN])
    build_log(sheets[SH_LOG])
    for market in ("미국", "한국"):
        build_sentiment(sheets[SH_SENT[market]], market)
        build_raw(sheets[SH_RAW[market]], f"원자료 — {market} 시장 심리", RAW_BLOCKS[market], RAW_HOWTO[market], lambda k: RAW_DAILY)
    build_macro(sheets[SH_MACRO])
    macro_blocks = [(k, f"{lab} — FRED {MACRO_RAW[k][3]}", MACRO_RAW[k][0], MACRO_RAW[k][1], MACRO_RAW[k][3]) for k, lab, *_ in MACRO_SERIES]
    build_raw(sheets[SH_MACRO_RAW], "원자료 — 거시 참고 패널", macro_blocks, MACRO_HOWTO, lambda k: MACRO_RAW[k][2])
    build_settings(sheets[SH_SET])

    sheets[SH_DASH].sheet_properties.tabColor = "4472C4"
    for name in (SH_RAW["미국"], SH_RAW["한국"], SH_MACRO_RAW):
        sheets[name].sheet_properties.tabColor = "A6A6A6"
    sheets[SH_SET].sheet_properties.tabColor = "7F7F7F"

    if data is not None:
        fill(sheets, data)
    wb.save(path)


def fill(sheets, data):
    """예시 데이터를 입력 칸에 쓴다."""
    for country, stocks in data["stocks"].items():
        ws = sheets[SH_STOCK[country]]
        ws["K3"] = data["price_date"]
        for r, s in enumerate(stocks, start=STOCK_FIRST):
            values = {
                "code": s.code, "name": s.name, "held": "Y" if s.held else None, "fin": "Y" if s.excluded_sector else None,
                "basis": s.basis, "ni": s.net_income, "eq": s.equity, "liab": s.liabilities, "op": s.op_income,
                "rev": s.revenue, "rev3": s.revenue_3y_ago, "mcap": s.market_cap, "debt": s.debt, "cash": s.cash,
                "p12": s.price_12m, "p1": s.price_1m,
            }
            for k, v in values.items():
                ws[f"{SC[k]}{r}"] = v
    ws = sheets[SH_HOLD]
    acc = data["account"]
    ws["D3"], ws["D4"], ws["D5"] = acc.total_krw, acc.cash_krw, acc.usdkrw
    for r, h in enumerate(data["holdings"], start=HOLD_FIRST):
        for c, v in zip("ABCDEFGH", (h.country, h.code, h.name, h.qty, h.avg_price, h.cost_krw, h.price, h.last_earnings)):
            ws[f"{c}{r}"] = v
    ws = sheets[SH_REASON]
    for r, x in enumerate(data["reasons"], start=REASON_FIRST):
        names = {s.code: s.name for stocks in data["stocks"].values() for s in stocks}
        row = (x.code, names.get(x.code, ""), x.no, x.text, "Y" if x.core else None, x.evidence or None, x.break_rule or None, x.checked, x.result or None)
        for c, v in zip("ABCDEFGHI", row):
            ws[f"{c}{r}"] = v
    ws = sheets[SH_LOG]
    for r, row in enumerate(data.get("log", []), start=5):
        for c, val in zip("ABCDEFGHIJK", row):
            ws[f"{c}{r}"] = val
    ws = sheets[SH_PLAN]
    for r, p in enumerate(data["plans"], start=PLAN_FIRST):
        for c, v in zip("ABCDE", (p.country, p.code, p.name, p.amount_krw, p.start)):
            ws[f"{c}{r}"] = v
    for market in ("미국", "한국"):
        ws = sheets[SH_RAW[market]]
        for key, _, dc, vc, _ in RAW_BLOCKS[market]:
            ws[f"{dc}2"], ws[f"{vc}2"] = "날짜", key
            for r, (d, v) in enumerate(data["sentiment"][market][key], start=RAW_FIRST):
                ws[f"{dc}{r}"], ws[f"{vc}{r}"] = d, v
    ws = sheets[SH_MACRO_RAW]
    for key, (dc, vc, _, fred) in MACRO_RAW.items():
        ws[f"{dc}2"], ws[f"{vc}2"] = "observation_date", fred
        for r, (d, v) in enumerate(data["macro"][key], start=RAW_FIRST):
            ws[f"{dc}{r}"], ws[f"{vc}{r}"] = d, v
